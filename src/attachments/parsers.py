"""Bounded local parsers that turn user files into typed temporary chunks."""
from __future__ import annotations

import csv
import hashlib
import io
import multiprocessing
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

import fitz
from PIL import Image, ImageOps

from src.attachments.detection import DetectedUpload
from src.attachments.models import (
    AttachmentChunk,
    AttachmentError,
    AttachmentLimits,
    AttachmentLocator,
    ParseCoverage,
)

_VISION_MAX_BYTES = 1_500_000
_VISION_MAX_LONG_EDGE = 2048
_VISION_FALLBACK_LONG_EDGE = 1600
_VISION_JPEG_QUALITIES = (88, 80, 72, 64)


@dataclass(frozen=True)
class ParseResult:
    chunks: tuple[AttachmentChunk, ...]
    parse_coverage: ParseCoverage
    parser_name: str
    parser_version: str
    warnings: tuple[str, ...] = ()
    normalized_image: bytes = b""
    normalized_media_type: str = ""
    normalized_width: int = 0
    normalized_height: int = 0


def _resize_long_edge(image: Image.Image, limit: int) -> Image.Image:
    width, height = image.size
    longest = max(width, height)
    if longest <= limit:
        return image.copy()
    scale = limit / longest
    return image.resize(
        (max(1, round(width * scale)), max(1, round(height * scale))),
        Image.Resampling.LANCZOS,
    )


def _png_bytes(image: Image.Image) -> bytes:
    if "A" in image.getbands() or "transparency" in image.info:
        image = image.convert("RGBA")
    else:
        image = image.convert("RGB")
    image.info.clear()
    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _jpeg_bytes(image: Image.Image, quality: int) -> bytes:
    if image.mode in {"RGBA", "LA"} or "transparency" in image.info:
        rgba = image.convert("RGBA")
        background = Image.new("RGB", rgba.size, "white")
        background.paste(rgba, mask=rgba.getchannel("A"))
        image = background
    else:
        image = image.convert("RGB")
    image.info.clear()
    output = io.BytesIO()
    image.save(
        output,
        format="JPEG",
        quality=quality,
        optimize=True,
        progressive=True,
        subsampling="4:2:0",
    )
    return output.getvalue()


def _encode_vision_image(
    image: Image.Image,
    *,
    max_bytes: int = _VISION_MAX_BYTES,
    max_long_edge: int = _VISION_MAX_LONG_EDGE,
    fallback_long_edge: int = _VISION_FALLBACK_LONG_EDGE,
    jpeg_qualities: tuple[int, ...] = _VISION_JPEG_QUALITIES,
) -> tuple[bytes, str, int, int]:
    primary = _resize_long_edge(image, max_long_edge)
    png = _png_bytes(primary)
    if len(png) <= max_bytes:
        return png, "image/png", primary.width, primary.height
    for quality in jpeg_qualities:
        jpeg = _jpeg_bytes(primary, quality)
        if len(jpeg) <= max_bytes:
            return jpeg, "image/jpeg", primary.width, primary.height
    fallback = _resize_long_edge(primary, fallback_long_edge)
    for quality in jpeg_qualities:
        jpeg = _jpeg_bytes(fallback, quality)
        if len(jpeg) <= max_bytes:
            return jpeg, "image/jpeg", fallback.width, fallback.height
    raise AttachmentError("attachment_image_normalization_failed")


def parse_attachment(
    path: Path,
    detected: DetectedUpload,
    limits: AttachmentLimits,
) -> ParseResult:
    if limits.parse_timeout_seconds <= 0:
        raise AttachmentError("attachment_parse_timeout")
    context = multiprocessing.get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    process = context.Process(
        target=_parse_worker,
        args=(str(path), detected, limits, sender),
        daemon=True,
    )
    process.start()
    sender.close()
    if not receiver.poll(limits.parse_timeout_seconds):
        process.terminate()
        process.join(timeout=1)
        receiver.close()
        raise AttachmentError("attachment_parse_timeout")
    try:
        item = receiver.recv()
    except EOFError as exc:
        raise AttachmentError("attachment_parse_failed") from exc
    finally:
        receiver.close()
    process.join(timeout=1)
    if process.is_alive():
        process.terminate()
        process.join(timeout=1)
    if item[0] == "error":
        raise AttachmentError(item[1], item[2])
    return item[1]


def _parse_worker(path: str, detected: DetectedUpload, limits: AttachmentLimits, output) -> None:
    try:
        output.send(("ok", _parse_direct(Path(path), detected, limits)))
    except AttachmentError as exc:
        output.send(("error", exc.code, str(exc)))
    except Exception as exc:
        output.send(("error", "attachment_parse_failed", type(exc).__name__))
    finally:
        output.close()


def _parse_direct(path: Path, detected: DetectedUpload, limits: AttachmentLimits) -> ParseResult:
    size = path.stat().st_size
    if detected.kind == "image":
        if size > limits.image_max_bytes:
            raise AttachmentError("attachment_limit_exceeded")
        return _parse_image(path, limits)
    if detected.kind == "pdf":
        if size > limits.office_max_bytes:
            raise AttachmentError("attachment_limit_exceeded")
        return _parse_pdf(path, limits)
    if detected.kind == "document":
        if size > limits.office_max_bytes:
            raise AttachmentError("attachment_limit_exceeded")
        return _parse_docx(path, limits)
    if detected.kind == "spreadsheet" and detected.extension == "xlsx":
        if size > limits.office_max_bytes:
            raise AttachmentError("attachment_limit_exceeded")
        return _parse_xlsx(path, limits)
    if detected.kind == "spreadsheet":
        if size > limits.text_max_bytes:
            raise AttachmentError("attachment_limit_exceeded")
        return _parse_csv(path)
    if size > limits.text_max_bytes:
        raise AttachmentError("attachment_limit_exceeded")
    return _parse_text(path)


def _parse_image(path: Path, limits: AttachmentLimits) -> ParseResult:
    chunks: tuple[AttachmentChunk, ...] = ()
    warnings: tuple[str, ...] = ()
    try:
        with Image.open(path) as probe:
            probe.verify()
        with Image.open(path) as image:
            if image.width * image.height > limits.image_max_pixels:
                raise AttachmentError("attachment_limit_exceeded")
            image.load()
            decoded = ImageOps.exif_transpose(image)
            decoded.load()
            normalized_bytes, normalized_type, width, height = _encode_vision_image(decoded)
            ocr_image = decoded.convert("RGB")
            try:
                import pytesseract

                ocr_data = pytesseract.image_to_data(
                    ocr_image, output_type=pytesseract.Output.DICT,
                )
                words = [
                    str(word).strip()
                    for word in ocr_data.get("text", [])
                    if str(word).strip()
                ]
                ocr_text = " ".join(words)[:4000]
                if ocr_text:
                    chunks = (_make_chunk(
                        ocr_text,
                        AttachmentLocator(region_label="ocr"),
                        "image-ocr",
                    ),)
            except Exception:
                warnings = ("ocr_unavailable",)
    except AttachmentError:
        raise
    except Exception as exc:
        raise AttachmentError("attachment_parse_failed") from exc
    return ParseResult(
        chunks=chunks,
        parse_coverage="partial" if chunks else "empty",
        parser_name="pillow",
        parser_version=getattr(Image, "__version__", "unknown"),
        warnings=warnings,
        normalized_image=normalized_bytes,
        normalized_media_type=normalized_type,
        normalized_width=width,
        normalized_height=height,
    )


def _parse_pdf(path: Path, limits: AttachmentLimits) -> ParseResult:
    document = fitz.open(path)
    try:
        if document.needs_pass:
            raise AttachmentError("attachment_parse_failed")
        if document.page_count > limits.pdf_max_pages:
            raise AttachmentError("attachment_limit_exceeded")
        chunks = []
        blank_pages = 0
        for index, page in enumerate(document):
            text = page.get_text("text").strip()
            if not text:
                blank_pages += 1
                continue
            chunks.extend(_bounded_chunks(text, AttachmentLocator(page=index + 1), f"pdf-{index + 1}"))
    finally:
        document.close()
    coverage: ParseCoverage = "empty" if not chunks else ("partial" if blank_pages else "full")
    return ParseResult(tuple(chunks), coverage, "pymupdf", fitz.VersionBind)


def _validate_ooxml(path: Path, limits: AttachmentLimits, expected_root: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000:
                raise AttachmentError("attachment_limit_exceeded")
            expanded = 0
            names = set()
            for entry in entries:
                pure = PurePosixPath(entry.filename)
                if pure.is_absolute() or ".." in pure.parts:
                    raise AttachmentError("attachment_parse_failed")
                expanded += entry.file_size
                if expanded > max(limits.office_max_bytes * 20, 64 * 1024 * 1024):
                    raise AttachmentError("attachment_limit_exceeded")
                names.add(entry.filename)
            if "[Content_Types].xml" not in names or not any(
                name.startswith(expected_root) for name in names
            ):
                raise AttachmentError("attachment_parse_failed")
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or b"vbaproject" in content_types:
                raise AttachmentError("unsupported_attachment_type")
    except AttachmentError:
        raise
    except (OSError, zipfile.BadZipFile, KeyError) as exc:
        raise AttachmentError("attachment_parse_failed") from exc


def _parse_docx(path: Path, limits: AttachmentLimits) -> ParseResult:
    from docx import Document
    from docx import __version__ as docx_version

    _validate_ooxml(path, limits, "word/")
    document = Document(path)
    chunks: list[AttachmentChunk] = []
    heading = ""
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            heading = text
        locator = AttachmentLocator(heading=heading or None, paragraph_range=f"{index}-{index}")
        chunks.extend(_bounded_chunks(text, locator, f"docx-p{index}"))
    for table_index, table in enumerate(document.tables, start=1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(row for row in rows if row.strip())
        if text:
            chunks.extend(_bounded_chunks(
                text,
                AttachmentLocator(heading=heading or None, table_index=table_index),
                f"docx-t{table_index}",
            ))
    return ParseResult(
        tuple(chunks),
        "full" if chunks else "empty",
        "python-docx",
        docx_version,
    )


def _parse_xlsx(path: Path, limits: AttachmentLimits) -> ParseResult:
    from openpyxl import __version__ as openpyxl_version
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    _validate_ooxml(path, limits, "xl/")
    # The upload service deliberately stages data under a generated ``.bin`` name;
    # pass a file object so openpyxl validates OOXML content rather than that safe
    # temporary suffix.
    source = path.open("rb")
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.sheetnames) > limits.spreadsheet_max_sheets:
            raise AttachmentError("attachment_limit_exceeded")
        chunks: list[AttachmentChunk] = []
        nonempty_total = 0
        for sheet in workbook.worksheets:
            values_by_coordinate: dict[tuple[int, int], str] = {}
            min_row = min_column = None
            max_row = max_column = 0
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None or str(cell.value) == "":
                        continue
                    nonempty_total += 1
                    if nonempty_total > limits.spreadsheet_max_nonempty_cells:
                        raise AttachmentError("attachment_limit_exceeded")
                    values_by_coordinate[(cell.row, cell.column)] = str(cell.value)
                    min_row = cell.row if min_row is None else min(min_row, cell.row)
                    min_column = cell.column if min_column is None else min(min_column, cell.column)
                    max_row = max(max_row, cell.row)
                    max_column = max(max_column, cell.column)
            if not values_by_coordinate or min_row is None or min_column is None:
                continue
            # Render only non-empty cells. Building the full bounding rectangle lets
            # two far-apart cells amplify into billions of empty strings.
            text = "\n".join(
                f"{get_column_letter(column)}{row}\t{value}"
                for (row, column), value in sorted(values_by_coordinate.items())
            )
            locator = AttachmentLocator(
                sheet=sheet.title,
                cell_range=(
                    f"{get_column_letter(min_column)}{min_row}:"
                    f"{get_column_letter(max_column)}{max_row}"
                ),
            )
            chunks.extend(_bounded_chunks(text, locator, f"xlsx-{sheet.title}"))
    finally:
        workbook.close()
        source.close()
    return ParseResult(
        tuple(chunks),
        "full" if chunks else "empty",
        "openpyxl",
        openpyxl_version,
    )


def _parse_csv(path: Path) -> ParseResult:
    text = _decode_text(path.read_bytes())
    rows = list(csv.reader(io.StringIO(text)))
    rendered = "\n".join("\t".join(row) for row in rows)
    width = max((len(row) for row in rows), default=0)
    locator = AttachmentLocator(sheet="CSV", cell_range=f"A1:{_column_name(width)}{len(rows)}")
    chunks = _bounded_chunks(rendered, locator, "csv") if rendered else []
    return ParseResult(tuple(chunks), "full" if chunks else "empty", "csv", "stdlib")


def _parse_text(path: Path) -> ParseResult:
    text = _decode_text(path.read_bytes())
    lines = text.splitlines()
    chunks: list[AttachmentChunk] = []
    buffered: list[str] = []
    buffer_start = 0

    def flush(end_line: int) -> None:
        nonlocal buffered, buffer_start
        if not buffered:
            return
        chunks.append(_make_chunk(
            "\n".join(buffered),
            AttachmentLocator(line_start=buffer_start, line_end=end_line),
            f"text-{buffer_start}-{end_line}",
        ))
        buffered = []
        buffer_start = 0

    for line_number, line in enumerate(lines, start=1):
        if len(line) > 4000:
            flush(line_number - 1)
            for segment_index, start in enumerate(range(0, len(line), 4000), start=1):
                chunks.append(_make_chunk(
                    line[start:start + 4000],
                    AttachmentLocator(line_start=line_number, line_end=line_number),
                    f"text-{line_number}-{segment_index}",
                ))
            continue
        projected = sum(len(item) for item in buffered) + len(buffered) + len(line)
        if buffered and projected > 4000:
            flush(line_number - 1)
        if not buffered:
            buffer_start = line_number
        buffered.append(line)
    flush(len(lines))
    return ParseResult(tuple(chunks), "full" if chunks else "empty", "text", "stdlib")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AttachmentError("attachment_parse_failed")


def _bounded_chunks(text: str, locator: AttachmentLocator, prefix: str) -> list[AttachmentChunk]:
    value = text.strip()
    if not value:
        return []
    chunks = []
    for index, start in enumerate(range(0, len(value), 4000), start=1):
        chunks.append(_make_chunk(value[start:start + 4000], locator, f"{prefix}-{index}"))
    return chunks


def _make_chunk(text: str, locator: AttachmentLocator, label: str) -> AttachmentChunk:
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    return AttachmentChunk(
        chunk_id=f"chunk-{hashlib.sha256(label.encode()).hexdigest()[:16]}",
        source_id="",
        locator=locator,
        text=text,
        content_sha256=digest,
        token_estimate=max(1, len(text) // 4),
    )


def _column_name(index: int) -> str:
    if index <= 0:
        return "A"
    value = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        value = chr(65 + remainder) + value
    return value
